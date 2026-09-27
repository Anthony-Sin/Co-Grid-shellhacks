"""export_data tool — write REAL overlap/project data to downloadable files.

Every row comes straight from the processed artifacts the public API
serves (overlaps.json / projects.geojson) through the same filters as
find_overlaps / list_projects — a generated file can never contain a
fabricated record (AGENTS.md §7). Files land in exports/ at the repo
root (gitignored generated output) and are served to the user at
/api/exports/<filename> — the tool returns that url so the model can
hand back a markdown download link.

Formats: csv (stdlib), xlsx (openpyxl), html (self-contained printable
report). An empty filtered result is an honest error — no empty file.
"""
from __future__ import annotations

import csv
import html as _html
import io
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from src.analysis.filters import filter_records
from src.processing.projection import scene_for_zone

from .tool_data import overlaps, projects

ROOT = Path(__file__).resolve().parents[2]
EXPORTS = ROOT / "exports"

_FORMATS = ("csv", "xlsx", "html")
_KINDS = ("overlaps", "projects")
_MAX_ROWS = 2000

# Same column set as /api/overlaps.csv — a generated workbook stays
# faithful to the spreadsheet the route serves.
OVERLAP_COLS = [
    "rank", "overlap_id", "project_a", "project_b", "utilities",
    "tier", "tier_label", "zone", "min_distance_km", "timeline_overlap",
    "timeline_adjacent", "shared_window_start", "shared_window_end",
    "closest_lon", "closest_lat", "score",
    "shared_row_km", "est_savings_low_usd", "est_savings_high_usd",
    "explanation", "map_link",
]

# project rows: record fields + a centroid so the sheet is GIS-friendly.
PROJECT_COLS = [
    "project_id", "utility", "name", "kind", "voltage_kv",
    "start_year", "end_year", "status", "location_confidence",
    "zones", "lon", "lat", "source",
]


def _overlap_rows(rows: list[dict]) -> list[list]:
    out = []
    for i, r in enumerate(rows, 1):  # rank = engine order, like the route
        win = r.get("shared_window") or {}
        cp = r.get("closest_point_a") or [None, None]
        cost = r.get("cost") or {}
        out.append([
            i, r.get("overlap_id"), r.get("project_a"), r.get("project_b"),
            "|".join(r.get("utilities") or []), r.get("tier"),
            r.get("tier_label"), r.get("zone"), r.get("min_distance_km"),
            r.get("timeline_overlap"), r.get("timeline_adjacent", False),
            win.get("start"), win.get("end"), cp[0], cp[1], r.get("score"),
            cost.get("shared_row_km"), cost.get("est_savings_usd_low"),
            cost.get("est_savings_usd_high"), r.get("explanation"),
            f"/?scene={scene_for_zone(r.get('zone'))}"
            f"&select={r.get('overlap_id')}&panel=0",
        ])
    return out


def _project_rows(feats: list[dict]) -> list[list]:
    out = []
    for f in feats:
        p = f.get("properties") or {}
        lon = lat = None
        geom = f.get("geometry")
        if geom:
            try:
                from shapely.geometry import shape
                c = shape(geom).centroid
                lon, lat = round(c.x, 6), round(c.y, 6)
            except Exception:  # geometry is decorative here — never fail a row
                pass
        out.append([
            p.get("project_id"), p.get("utility"), p.get("name"),
            p.get("kind"), p.get("voltage_kv"), p.get("start_year"),
            p.get("end_year"), p.get("status"), p.get("location_confidence"),
            "|".join(p.get("zones") or []), lon, lat, p.get("source"),
        ])
    return out


def _write_csv(path: Path, header: list[str], rows: list[list]) -> None:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    path.write_text(buf.getvalue())


def _write_xlsx(path: Path, title: str, header: list[str],
                rows: list[list]) -> dict | None:
    """openpyxl workbook; returns an error dict (not a crash) if the
    optional dep is missing."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return {"ok": False,
                "error": "xlsx export needs openpyxl — "
                         "`./venv/bin/pip install openpyxl`"}
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]  # excel sheet-name cap
    ws.append(header)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = Font(bold=True)
    for idx in range(1, len(header) + 1):
        width = max((len(str(r[idx - 1])) for r in [header, *rows[:50]]),
                    default=8)
        ws.column_dimensions[get_column_letter(idx)].width = \
            min(max(width + 2, 8), 60)
    wb.save(path)
    return None


def _write_html(path: Path, title: str, header: list[str], rows: list[list],
                generated: str, filter_summary: str) -> None:
    """Minimal self-contained printable report — inline CSS only, no
    scripts, every cell html-escaped."""
    esc = _html.escape
    thead = "".join(f"<th>{esc(str(h))}</th>" for h in header)
    body = "\n".join(
        "<tr>" + "".join(
            f"<td>{esc('' if v is None else str(v))}</td>" for v in row)
        + "</tr>"
        for row in rows)
    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)}</title>
<style>
body{{font:13px/1.45 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;
  color:#1c1c1c;margin:32px;}}
h1{{font-size:18px;margin:0 0 4px;}}
.meta{{color:#666;margin:0 0 16px;font-size:12px;}}
table{{border-collapse:collapse;width:100%;}}
th,td{{border:1px solid #ccc;padding:4px 8px;text-align:left;
  vertical-align:top;font-size:11px;}}
th{{background:#f0ead8;position:sticky;top:0;}}
tr:nth-child(even) td{{background:#faf8f2;}}
@media print{{body{{margin:8mm;}}th{{position:static;}}}}
</style></head><body>
<h1>{esc(title)}</h1>
<p class="meta">generated {esc(generated)} &middot; {esc(filter_summary)}
&middot; {len(rows)} row{"s" if len(rows) != 1 else ""}
&middot; real filed data only (no CEII)</p>
<table><thead><tr>{thead}</tr></thead>
<tbody>
{body}
</tbody></table>
</body></html>
"""
    path.write_text(doc)


def _unique_path(kind: str, ext: str, stamp: str) -> Path:
    """co-grid-<kind>-<YYYYMMDD-HHMMSS>.<ext>; bump -2/-3 if a same-second
    export already claimed the name instead of silently overwriting."""
    base = EXPORTS / f"co-grid-{kind}-{stamp}"
    path = base.with_suffix(f".{ext}")
    n = 2
    while path.exists():
        path = Path(f"{base}-{n}.{ext}")
        n += 1
    return path


def tool_export_data(format: str = "csv",
                     kind: str = "overlaps",
                     utility: str | None = None,
                     utilities: str | list | None = None,
                     tier: int | None = None,
                     zone: str | None = None,
                     source: str | None = None,
                     timeline_only: bool = False,
                     adjacent_only: bool = False,
                     limit: int = 200) -> dict:
    """Write filtered REAL records to a downloadable file under exports/.

    Overlap exports accept the full find_overlaps filter set (utility
    either-side, utilities exact pair, tier, zone substring,
    timeline_only / adjacent_only); project exports accept the
    list_projects filters (utility, zone, source substring). `limit`
    caps rows at 2000 — the file is the FULL filtered set up to that
    cap, not the preview window the chat tools show.
    """
    fmt = str(format or "csv").strip().lower().lstrip(".")
    if fmt not in _FORMATS:
        return {"ok": False,
                "error": f"unknown format {format!r} — use csv|xlsx|html"}
    k = str(kind or "overlaps").strip().lower()
    if k not in _KINDS:
        return {"ok": False,
                "error": f"unknown kind {kind!r} — use overlaps|projects"}
    lim = max(1, min(int(limit or 200), _MAX_ROWS))

    applied: dict[str, Any] = {}
    if k == "overlaps":
        try:
            rows = filter_records(
                overlaps(), utility=utility, utilities=utilities,
                tier=tier, zone=zone, timeline_only=timeline_only,
                adjacent_only=adjacent_only)
        except ValueError as e:
            return {"ok": False, "error": str(e)}
        applied = {x: y for x, y in {
            "utility": utility, "utilities": utilities, "tier": tier,
            "zone": zone, "timeline_only": timeline_only or None,
            "adjacent_only": adjacent_only or None}.items()
            if y is not None}
        header, total = OVERLAP_COLS, len(rows)
        data_rows = _overlap_rows(rows[:lim])
        title = "CO-GRID — coordination overlap export"
    else:
        feats = projects()
        if utility:
            feats = [f for f in feats
                     if f["properties"].get("utility") == utility]
        if zone:
            feats = [f for f in feats
                     if zone in (f["properties"].get("zones") or [])]
        if source:
            s = str(source).strip().lower()
            feats = [f for f in feats
                     if s in (f["properties"].get("source") or "").lower()]
        applied = {x: y for x, y in
                   {"utility": utility, "zone": zone,
                    "source": source}.items() if y is not None}
        ignored = [x for x, y in {"tier": tier, "utilities": utilities,
                                  "timeline_only": timeline_only,
                                  "adjacent_only": adjacent_only}.items()
                   if y]
        if ignored:
            applied["note_ignored_for_projects"] = ignored
        header, total = PROJECT_COLS, len(feats)
        data_rows = _project_rows(feats[:lim])
        title = "CO-GRID — planned project export"

    if not data_rows:
        return {"ok": False,
                "error": f"0 {k} match the given filters — nothing to "
                         "export (the file would be empty; loosen the "
                         "filters or check stats.zones_available)",
                "filters": applied}
    EXPORTS.mkdir(exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    path = _unique_path(k, fmt, stamp)
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    filter_summary = ("filters: " + ", ".join(f"{a}={b}" for a, b in
                                              applied.items())
                      if applied else "filters: none")

    if fmt == "csv":
        _write_csv(path, header, data_rows)
    elif fmt == "xlsx":
        err = _write_xlsx(path, title, header, data_rows)
        if err:
            return err
    else:
        _write_html(path, title, header, data_rows, generated,
                    filter_summary)

    return {
        "ok": True,
        "filename": path.name,
        "kind": k,
        "format": fmt,
        "rows": len(data_rows),
        "total_matching": total,
        "capped": total > lim,
        "filters": applied,
        "url": f"/api/exports/{path.name}",
        "note": "share the markdown link with the user",
    }
